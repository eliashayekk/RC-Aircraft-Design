from pathlib import Path
import pandas as pd
import re

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference


# -------------------------------------------------
# SETTINGS
# -------------------------------------------------

data_folder = Path("data")

# Find all exported XFLR5 .txt files
files = list(data_folder.glob("*.txt"))

results = []
polar_data = {}


# -------------------------------------------------
# PROCESS EACH FILE
# -------------------------------------------------

for file in files:

    print(f"\nProcessing: {file.name}")

    # -------------------------------------------------
    # 1. READ FILE
    # -------------------------------------------------

    with open(file, "r") as f:
        lines = f.readlines()

    # -------------------------------------------------
    # 2. FIND TABLE HEADER
    # -------------------------------------------------

    header_index = None

    for i, line in enumerate(lines):

        line_lower = line.lower()

        if (
            "alpha" in line_lower
            and "cl" in line_lower
            and "cd" in line_lower
        ):
            header_index = i
            break

    if header_index is None:
        print("  ERROR: Could not find table header")
        continue

    print(f"  Header found at line {header_index + 1}")

    # -------------------------------------------------
    # 3. READ COLUMN NAMES
    # -------------------------------------------------

    header = lines[header_index].split()

    print("  Columns found:", header)

    # -------------------------------------------------
    # 4. READ NUMERICAL DATA
    # -------------------------------------------------

    rows = []

    for line in lines[header_index + 1:]:

        stripped = line.strip()

        if not stripped:
            continue

        parts = stripped.split()

        if len(parts) != len(header):
            continue

        try:
            values = [float(value) for value in parts]
            rows.append(values)

        except ValueError:
            continue

    if not rows:
        print("  ERROR: No numerical data found")
        continue

    # -------------------------------------------------
    # 5. CREATE DATAFRAME
    # -------------------------------------------------

    df = pd.DataFrame(
        rows,
        columns=header
    )

    print(f"  Loaded {len(df)} operating points")

    # -------------------------------------------------
    # 6. CHECK REQUIRED COLUMNS
    # -------------------------------------------------

    required_columns = [
        "alpha",
        "CL",
        "CD"
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing_columns:
        print(
            f"  ERROR: Missing columns: {missing_columns}"
        )
        continue

    # -------------------------------------------------
    # 7. CALCULATE CL MAX
    # -------------------------------------------------

    cl_max_index = df["CL"].idxmax()

    cl_max = df.loc[
        cl_max_index,
        "CL"
    ]

    alpha_stall = df.loc[
        cl_max_index,
        "alpha"
    ]

    # -------------------------------------------------
    # 8. CALCULATE L/D
    # -------------------------------------------------

    # Create L/D column in the main dataframe
    df["L_D"] = None

    valid_cd = df["CD"] > 0

    df.loc[
        valid_cd,
        "L_D"
    ] = (
        df.loc[valid_cd, "CL"]
        / df.loc[valid_cd, "CD"]
    )

    # Create clean dataframe for L/D analysis
    valid_df = df[
        df["CD"] > 0
    ].copy()

    if valid_df.empty:
        print("  ERROR: No valid CD values")
        continue

    ld_max_index = valid_df["L_D"].idxmax()

    ld_max = valid_df.loc[
        ld_max_index,
        "L_D"
    ]

    alpha_ld_max = valid_df.loc[
        ld_max_index,
        "alpha"
    ]

    # -------------------------------------------------
    # 9. PITCHING MOMENT AT MAX L/D
    # -------------------------------------------------

    if "Cm" in valid_df.columns:

        cm_at_ld_max = valid_df.loc[
            ld_max_index,
            "Cm"
        ]

    elif "CM" in valid_df.columns:

        cm_at_ld_max = valid_df.loc[
            ld_max_index,
            "CM"
        ]

    else:

        cm_at_ld_max = None

    # -------------------------------------------------
    # 10. READ AIRFOIL NAME
    # -------------------------------------------------

    filename = file.stem

    if "_T1" in filename:

        airfoil = filename.split("_T1")[0]

    elif "_Re" in filename:

        airfoil = filename.split("_Re")[0]

    else:

        airfoil = filename

    # -------------------------------------------------
    # 11. READ REYNOLDS NUMBER
    #
    # Example:
    # Re0.077 = 77,000
    # -------------------------------------------------

    match = re.search(
        r"_Re([0-9.]+)_",
        filename
    )

    if match:

        re_millions = float(
            match.group(1)
        )

        reynolds = int(
            re_millions * 1_000_000
        )

    else:

        reynolds = None

    # -------------------------------------------------
    # 12. STORE FULL POLAR DATA
    # -------------------------------------------------

    polar_key = f"{airfoil}_{reynolds}"

    polar_data[
        polar_key
    ] = df.copy()

    # -------------------------------------------------
    # 13. STORE SUMMARY RESULTS
    # -------------------------------------------------

    results.append(
        {
            "Airfoil": airfoil,
            "Re": reynolds,
            "CL_max": cl_max,
            "alpha_stall_deg": alpha_stall,
            "LD_max": ld_max,
            "alpha_LD_max_deg": alpha_ld_max,
            "Cm_at_LD_max": cm_at_ld_max
        }
    )


# -------------------------------------------------
# CREATE FINAL SUMMARY
# -------------------------------------------------

if not results:

    print("\nNo valid polar files were processed.")

else:

    summary = pd.DataFrame(results)

    summary = summary.sort_values(
        by=[
            "Airfoil",
            "Re"
        ],
        na_position="last"
    )

    # -------------------------------------------------
    # CLEAN DISPLAY NAMES
    # -------------------------------------------------

    summary["Airfoil"] = (
        summary["Airfoil"]
        .str.upper()
    )

    summary = summary.rename(
        columns={
            "CL_max": "CL max",
            "alpha_stall_deg": "Stall angle (deg)",
            "LD_max": "L/D max",
            "alpha_LD_max_deg": "Angle at max L/D (deg)",
            "Cm_at_LD_max": "Cm at max L/D"
        }
    )

    summary = summary.round(
        {
            "CL max": 4,
            "Stall angle (deg)": 2,
            "L/D max": 2,
            "Angle at max L/D (deg)": 2,
            "Cm at max L/D": 4
        }
    )

    # -------------------------------------------------
    # PRINT SUMMARY
    # -------------------------------------------------

    print("\n")
    print("=" * 100)
    print("AIRFOIL POLAR SUMMARY")
    print("=" * 100)

    print(
        summary.to_string(
            index=False
        )
    )

    # -------------------------------------------------
    # SAVE CSV
    # -------------------------------------------------

    summary.to_csv(
        "airfoil_summary.csv",
        index=False
    )

    # -------------------------------------------------
    # SAVE EXCEL WORKBOOK
    # -------------------------------------------------

    excel_file = "airfoil_analysis.xlsx"

    with pd.ExcelWriter(
        excel_file,
        engine="openpyxl"
    ) as writer:

        # Summary sheet
        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False
        )

        # Raw polar sheets
        for polar_name, polar_df in polar_data.items():

            sheet_name = polar_name[:31]

            polar_df.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False
            )

    # -------------------------------------------------
    # FORMAT WORKBOOK
    # -------------------------------------------------

    workbook = load_workbook(
        excel_file
    )

    for worksheet in workbook.worksheets:

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # Header styling
        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Center body cells
        for row in worksheet.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.alignment = Alignment(
                    horizontal="center"
                )

        # Auto-size columns
        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = (
                column_cells[0]
                .column_letter
            )

            for cell in column_cells:

                if cell.value is not None:

                    length = len(
                        str(cell.value)
                    )

                    if length > max_length:
                        max_length = length

            worksheet.column_dimensions[
                column_letter
            ].width = max_length + 3

    # -------------------------------------------------
    # CREATE CHARTS
    # -------------------------------------------------

    for sheet_name in workbook.sheetnames:

        if sheet_name == "Summary":
            continue

        ws = workbook[
            sheet_name
        ]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        # ---------------------------------------------
        # CL vs ALPHA
        # ---------------------------------------------

        if (
            "alpha" in headers
            and "CL" in headers
        ):

            chart = LineChart()

            chart.title = "CL vs Angle of Attack"
            chart.x_axis.title = "Angle of Attack (deg)"
            chart.y_axis.title = "CL"

            data = Reference(
                ws,
                min_col=headers["CL"],
                min_row=1,
                max_row=ws.max_row
            )

            categories = Reference(
                ws,
                min_col=headers["alpha"],
                min_row=2,
                max_row=ws.max_row
            )

            chart.add_data(
                data,
                titles_from_data=True
            )

            chart.set_categories(
                categories
            )

            chart.height = 8
            chart.width = 14

            ws.add_chart(
                chart,
                "N2"
            )

        # ---------------------------------------------
        # CD vs ALPHA
        # ---------------------------------------------

        if (
            "alpha" in headers
            and "CD" in headers
        ):

            chart = LineChart()

            chart.title = "CD vs Angle of Attack"
            chart.x_axis.title = "Angle of Attack (deg)"
            chart.y_axis.title = "CD"

            data = Reference(
                ws,
                min_col=headers["CD"],
                min_row=1,
                max_row=ws.max_row
            )

            categories = Reference(
                ws,
                min_col=headers["alpha"],
                min_row=2,
                max_row=ws.max_row
            )

            chart.add_data(
                data,
                titles_from_data=True
            )

            chart.set_categories(
                categories
            )

            chart.height = 8
            chart.width = 14

            ws.add_chart(
                chart,
                "N18"
            )

        # ---------------------------------------------
        # L/D vs ALPHA
        # ---------------------------------------------

        if (
            "alpha" in headers
            and "L_D" in headers
        ):

            chart = LineChart()

            chart.title = "L/D vs Angle of Attack"
            chart.x_axis.title = "Angle of Attack (deg)"
            chart.y_axis.title = "L/D"

            data = Reference(
                ws,
                min_col=headers["L_D"],
                min_row=1,
                max_row=ws.max_row
            )

            categories = Reference(
                ws,
                min_col=headers["alpha"],
                min_row=2,
                max_row=ws.max_row
            )

            chart.add_data(
                data,
                titles_from_data=True
            )

            chart.set_categories(
                categories
            )

            chart.height = 8
            chart.width = 14

            ws.add_chart(
                chart,
                "N34"
            )

        # ---------------------------------------------
        # Cm vs ALPHA
        # ---------------------------------------------

        cm_column = None

        if "Cm" in headers:
            cm_column = "Cm"

        elif "CM" in headers:
            cm_column = "CM"

        if (
            "alpha" in headers
            and cm_column
        ):

            chart = LineChart()

            chart.title = "Cm vs Angle of Attack"
            chart.x_axis.title = "Angle of Attack (deg)"
            chart.y_axis.title = "Cm"

            data = Reference(
                ws,
                min_col=headers[cm_column],
                min_row=1,
                max_row=ws.max_row
            )

            categories = Reference(
                ws,
                min_col=headers["alpha"],
                min_row=2,
                max_row=ws.max_row
            )

            chart.add_data(
                data,
                titles_from_data=True
            )

            chart.set_categories(
                categories
            )

            chart.height = 8
            chart.width = 14

            ws.add_chart(
                chart,
                "N50"
            )

    # -------------------------------------------------
    # SAVE FINAL WORKBOOK
    # -------------------------------------------------

    workbook.save(
        excel_file
    )

    print("\nFiles created:")
    print("  airfoil_summary.csv")
    print("  airfoil_analysis.xlsx")